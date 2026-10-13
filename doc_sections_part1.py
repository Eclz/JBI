"""
Part 1: Cover, Executive Summary, Stakeholders, Multi-Role Deep Dive
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from doc_builder_part1 import (
    COLOR_NAVY, COLOR_BLUE, COLOR_SLATE, COLOR_BODY, COLOR_MUTED, COLOR_AMBER, COLOR_GREEN,
    add_heading_1, add_heading_2, add_heading_3, add_heading_4,
    add_paragraph, add_bullet, add_callout, add_image_with_caption, add_styled_table
)

def build_part_1(doc):
    print("Building Part 1: Cover, Board Executive Summary, Stakeholders, and Multi-Role Architecture...")

    # ==========================================
    # COVER / TITLE BLOCK
    # ==========================================
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("JBI UNIVERSITY MANAGEMENT SYSTEM")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Master Software Engineering Architecture, Formal Database Normalization & Institutional Operating Manual")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_BLUE

    # Meta Table
    meta_headers = ["Document Parameter", "Specification Details"]
    meta_data = [
        ["System Reference", "JBI Tertiary Education Enterprise ERP (Web Portal & Administrative Engine)"],
        ["Target Audience", "University Board of Directors, Vice-Chancellor's Office, Future System Architects & Software Engineers"],
        ["Version & Release", "Version 2.2 Enterprise Production Release (October 2026)"],
        ["Core Framework Stack", "Laravel 11.x, PHP 8.2.12+, MySQL 8.x InnoDB, Bootstrap 5.3, Vite/JS"],
        ["Multi-Role RBAC Model", "Relational Many-to-Many Pivot (role_user) with Dynamic Academic Leadership Resolution"],
        ["Excluded Stakeholder", "Parent / Guardian (strictly excluded from models, interfaces, and documentation per institutional mandate)"],
        ["Author & Reviewer", "Lead Software Engineer & Principal System Analyst"]
    ]
    add_styled_table(doc, meta_headers, meta_data, [Inches(2.2), Inches(4.3)])

    doc.add_page_break()

    # ==========================================
    # SECTION 1: EXECUTIVE SUMMARY FOR THE UNIVERSITY BOARD
    # ==========================================
    add_heading_1(doc, "1. Executive Summary for the University Board of Directors")
    add_paragraph(doc, "This comprehensive engineering and system architecture specification provides the University Board of Directors, the Vice-Chancellor, and institutional governing bodies with a strategic and technical appraisal of the JBI University Management System. Designed specifically for contemporary higher education governance, the platform unifies student admissions, fee reconciliation, academic operations, examinations, online learning, student e-voting democracy, and human capital administration into an integrated digital campus.")

    add_heading_2(doc, "1.1. Strategic Institutional Value & Return on Investment (ROI)")
    add_paragraph(doc, "Higher education institutions face unprecedented challenges in maintaining academic integrity, optimizing operating costs, and eliminating revenue leakages. The JBI platform resolves these strategic challenges through five pillars of automation:")
    add_bullet(doc, "Replaces manual paper applications and fragmented departmental spreadsheets with a centralized, auditable admission pipeline. Prospective applicants submit credentials digitally, undergo automated entry qualification verification, and receive instant admission letters and activation credentials.", "1. Centralized Admissions Velocity: ");
    add_bullet(doc, "Eliminates cash-handling risks through integrated Payment Reference Number (PRN) generation. All fee records are reconciled directly against bank gateway transactions, establishing verifiable audit trails and preventing revenue leakage.", "2. Financial Anti-Fraud Integrity: ");
    add_bullet(doc, "Protects institutional accreditation by automating prerequisite enforcement, course load restrictions, continuous assessment moderation, and cumulative Grade Point Average (CGPA) calculations.", "3. Academic Rigor Automation: ");
    add_bullet(doc, "Empowers senior faculty members to serve concurrently across instructional, departmental (HOD), and divisional (Dean) roles under a single secure identity, ending administrative fragmentation.", "4. Multi-Role Institutional Agility: ");
    add_bullet(doc, "Provides a transparent, tamper-resistant digital voting platform for student union elections, enforcing an automated academic vetting rule that immediately disqualifies students with active course retakes or failing grades.", "5. Democratic Campus Governance: ");

    add_heading_2(doc, "1.2. Fiduciary Governance, Auditability & Regulatory Compliance")
    add_paragraph(doc, "The JBI platform incorporates end-to-end auditability across every financial, academic, and administrative interaction. Every grade alteration, invoice adjustment, fee waiver, and student status modification is committed to immutable audit logs containing timestamps, originating IP addresses, user identity tokens, and before/after state snapshots. This guarantees full compliance with tertiary education accreditation bodies, external auditing standards, and national data protection regulations.")

    add_callout(doc, "Executive Institutional Signoff & Security Guarantee", 
        "The system enforces strict principle-of-least-privilege authorization. Even in multi-role configurations, administrative powers are tightly isolated behind permission gates. Central administrative resets (/clear, /generate) are fortified behind multi-factor authentication and SuperAdmin authorization checks. The institution retains complete sovereign data ownership with zero dependency on proprietary vendor lock-in.",
        bg_hex='EFF6FF', border_hex='2563EB', title_color=COLOR_BLUE)

    # ==========================================
    # SECTION 2: STAKEHOLDER & UNIVERSITY ACTOR TAXONOMY
    # ==========================================
    add_heading_1(doc, "2. Institutional Stakeholder & Actor Taxonomy (Parent Excluded)")
    add_paragraph(doc, "In accordance with institutional guidelines, all university actors have been formally modeled and assigned precise functional boundaries. As explicitly instructed, the Parent/Guardian role is completely omitted from this documentation and the university system access model.")

    add_heading_2(doc, "2.1. In-Depth Stakeholder Profiles & Clearance Hierarchy")

    add_heading_3(doc, "1. Super Administrator & Central IT Staff")
    add_paragraph(doc, "The Central IT Directorate exercises highest-level operational responsibility for system availability, network security, database backups, and disaster recovery.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Manages system configurations, executes database migrations, provisions official staff accounts, manages the role and permission catalog, oversees queue workers, and inspects security telemetry and error logs.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 5 (Unrestricted System Root with Gate::before universal bypass privileges).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "2. University Board & Vice-Chancellor's Office")
    add_paragraph(doc, "The executive governing body presiding over institutional strategy, financial sustainability, academic quality, and statutory compliance.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Monitors institutional KPIs, reviews cross-faculty enrollment trends, audits multi-currency fee collection revenues, reviews faculty budget allocations, and approves strategic expansions.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 4 (Executive Audit & Read-Only Reporting across all institutional divisions).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "3. Registrar & Admissions Directorate")
    add_paragraph(doc, "The central custodian of student academic records, institutional matriculation, degree awards, and formal university admissions.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Reviews incoming prospective student applications, verifies uploaded qualification certificates, approves or rejects applications, issues official matriculation numbers (ADM-YYYY-XXXX), certifies semester graduation lists, and seals academic transcripts.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 3 (Registry Authority with full CRUD over admissions and student records).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "4. Bursar & Chief Financial Officer")
    add_paragraph(doc, "The chief financial officer responsible for tuition fee structures, billing ledgers, banking reconciliations, and payroll disbursement.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Configures academic year fee schedules, generates student fee invoices, validates PRN payment clearances, tracks bank transactions, administers staff payroll, and generates institutional financial statements.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 3 (Financial Authority over ledgers, fee structures, and accounting entries).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "5. Dean of Faculty")
    add_paragraph(doc, "Senior academic executive leading a major academic division (e.g. Faculty of Technology, Faculty of Business & Economics).", bold_prefix="Mandate: ")
    add_paragraph(doc, "Oversees academic quality across departments, reviews and endorses moderated semester grades, approves departmental budget allocations, and leads international accreditation and corporate partnerships.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 3 (Divisional Academic Leadership with Dean Administration permissions).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "6. Head of Department (HOD)")
    add_paragraph(doc, "Academic leader responsible for instructional delivery, curriculum coherence, and staff workload within an academic department.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Allocates course teaching assignments to department lecturers, reviews syllabi, moderates continuous assessment and examination papers, and audits student degree progression.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 3 (Departmental Academic Leadership with course and lecturer moderation rights).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "7. Academic Lecturer")
    add_paragraph(doc, "Teaching and research faculty directly responsible for classroom instruction, student mentoring, and coursework assessment.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Uploads course syllabus materials, creates assignments and online quizzes, records attendance via dynamic daily QR codes, grades submissions, and submits continuous assessment scores.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 2 (Academic Operational over assigned courses and enrolled students).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "8. Student Body (Enrolled Students)")
    add_paragraph(doc, "Matriculated undergraduate and postgraduate scholars pursuing accredited academic degrees.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Enrolls in semester courses, accesses digital lecture materials, submits assignments, attempts quizzes and examinations, reviews semester GPA and transcripts, tracks tuition balances, submits support tickets, and votes in elections.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 1 (Student Portal access limited to personal records, enrolled courses, and voting).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "9. Estates & Facilities Manager")
    add_paragraph(doc, "Administrative authority overseeing the physical campus infrastructure, lecture theaters, laboratories, and residence halls.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Manages classroom allocations, coordinates campus maintenance tickets, schedules facility bookings, and manages student hostel room requests.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 2 (Facilities Operational over physical assets, rooms, and halls).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "10. Human Resources (HR) Manager")
    add_paragraph(doc, "Administrative authority responsible for staff recruitment, employee files, contracts, leave, and payroll profiles.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Maintains staff directories, posts academic and non-academic vacancies, tracks recruitment pipelines, and manages employee contracts and leave allowances.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 3 (HR Authority over staff records and recruitment pipelines).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "11. University Librarian & Circulation Staff")
    add_paragraph(doc, "Information professionals curating physical and digital learning repositories for teaching and research.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Catalogues books, e-journals, and monographs, manages patron loans, issues overdue fine notices, and manages physical reservations.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 2 (Library Operational over catalog and circulation loans).", bold_prefix="Clearance Level: ")

    add_heading_3(doc, "12. Electoral Commission Trustee")
    add_paragraph(doc, "Independent constitutional body overseeing transparent and fair student leadership elections.", bold_prefix="Mandate: ")
    add_paragraph(doc, "Schedules election seasons, vets candidate portfolios, verifies academic clearance, oversees secret electronic balloting, and certifies live vote tallies.", bold_prefix="Key Responsibilities: ")
    add_paragraph(doc, "Level 3 (Electoral Governance over voting sessions, candidate vetting, and tallies).", bold_prefix="Clearance Level: ")

    # Stakeholder-to-Subsystem Matrix
    add_heading_2(doc, "2.2. Master Stakeholder-to-Subsystem Matrix")
    matrix_headers = ["University Actor", "Admissions", "Academics", "LMS / Exam", "Finance", "E-Voting", "Estates/HR"]
    matrix_data = [
        ["Super Administrator", "Full Access", "Full Access", "Full Access", "Full Access", "Full Access", "Full Access"],
        ["Board & VC", "Audit / View", "Audit / View", "Audit / View", "Audit / View", "Audit / View", "Audit / View"],
        ["Registrar", "Manage / Seal", "Manage / Seal", "View Records", "View Invoices", "Certify Roll", "View Directory"],
        ["Bursar / Finance", "Fee Clearance", "View Programs", "Exam Clearance", "Manage / Reconcile", "Fee Clearance", "Manage Payroll"],
        ["Dean of Faculty", "Faculty Queue", "Strategic Quality", "Endorse Grades", "Faculty Budget", "Oversee Faculty", "Faculty Space"],
        ["Head of Dept (HOD)", "Dept Inquiries", "Assign Courses", "Moderate Grades", "Dept Budget", "Dept Candidates", "Dept Facilities"],
        ["Academic Lecturer", "No Access", "Assigned Courses", "Teach / Grade", "No Access", "No Access", "Lab Requests"],
        ["Student Body", "Apply / Status", "Enroll Courses", "Submit / Quiz", "Pay / PRN", "Candidacy / Vote", "Hostel Request"],
        ["Estates Manager", "No Access", "Room Timetable", "Exam Halls", "Budget Assets", "Polling Venues", "Manage All"],
        ["HR Manager", "No Access", "Staff Allocation", "Faculty Record", "Payroll Sync", "Staff Oversight", "Manage Staff"],
        ["Librarian", "No Access", "Course Reading", "Material Sync", "Library Fines", "No Access", "Library Space"],
        ["Electoral Trustee", "No Access", "Vetting Roll", "Retake Vetting", "No Access", "Manage Ballot", "Polling Stations"]
    ]
    add_styled_table(doc, matrix_headers, matrix_data, [Inches(1.5), Inches(0.8), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.9)])

    # ==========================================
    # SECTION 3: MULTI-ROLE ARCHITECTURAL DEEP DIVE
    # ==========================================
    add_heading_1(doc, "3. Multi-Role & Dynamic Appointment Architecture")
    add_paragraph(doc, "In tertiary education administration, faculty and staff members frequently hold multiple concurrent responsibilities. A single user account must seamlessly support these multi-role appointments without compromising security or requiring multiple sets of credentials.")

    add_heading_2(doc, "3.1. Architectural Problem Analysis")
    add_paragraph(doc, "Conventional single-role ERP systems enforce rigid schemas (e.g. users.role = 'dean'). This causes severe operational breakdowns in university environments:")
    add_bullet(doc, "When a professor is appointed Dean, their role changes to 'dean'. In single-role systems, this causes them to lose access to the Lecturer portal, preventing them from uploading materials, grading assignments, or recording QR attendance for courses they continue to teach.", "Problem 1 (Loss of Teaching Interface): ");
    add_bullet(doc, "Department Heads (HODs) who also serve as admissions review committee members or research directors are forced to maintain multiple email addresses and passwords, leading to security fatigue and fragmented audit logs.", "Problem 2 (Credential Fragmentation): ");
    add_bullet(doc, "Staff members in Bursary or Estates who teach part-time or adjunct evening courses cannot access learning management systems without granting full administrative privileges to the entire institution.", "Problem 3 (Over-Privileged Workarounds): ");

    add_heading_2(doc, "3.2. Technical Solution: Relational Many-to-Many RBAC Engine")
    add_paragraph(doc, "The JBI architecture was enhanced to implement a unified, multi-layered role and permission evaluation pipeline consisting of four interrelated components:")
    add_bullet(doc, "A dedicated pivot table (role_user) links user accounts to multiple role definitions in the roles catalog. It includes an is_primary boolean flag ensuring the user retains a default dashboard context upon authentication.", "1. Relational Associative Pivot: ");
    add_bullet(doc, "The User model dynamically evaluates academic leadership appointments in faculties.dean_id and departments.head_of_department_id. Faculty members heading a division dynamically inherit Dean/HOD permissions without requiring manual database overrides.", "2. Dynamic Leadership Evaluator: ");
    add_bullet(doc, "The hasPermission($module, $action) engine aggregates permissions across all active assigned roles. If any role possessed by the user grants an authorization, access is approved.", "3. Composite Permission Aggregator: ");
    add_bullet(doc, "The administrative user management interfaces (create and edit forms) provide an interactive multi-role selector allowing administrators to assign secondary concurrent roles effortlessly.", "4. Multi-Role Admin Management UI: ");

    add_heading_2(doc, "3.3. Multi-Role Verification & Test Harness Results")
    add_paragraph(doc, "The multi-role engine was rigorously tested against four real-world university scenarios using automated test harnesses:")

    test_headers = ["Test Scenario", "Assigned Roles", "Dynamic Identifiers", "Permission Assertions", "Operational Verification Result"]
    test_data = [
        ["Scenario 1:\nLecturer + Dean", "Primary: Lecturer\nSecondary: Dean", "isFaculty() = true\nisDean() = true", "dean_administration.view: PASS\ncourses.view: PASS", "100% SUCCESS: User can upload lecture slides and grade exams in Lecturer portal, while simultaneously accessing /faculty/dean/budget and quality dashboards."],
        ["Scenario 2:\nLecturer + HOD", "Primary: Lecturer\nSecondary: Head of Dept", "isFaculty() = true\nisHeadOfDepartment() = true", "department_courses.view: PASS\ncourses.view: PASS", "100% SUCCESS: User retains teaching duties in Computer Science, while exercising moderation and workload allocation over departmental faculty."],
        ["Scenario 3:\nFinance + Lecturer", "Primary: Finance Officer\nSecondary: Lecturer", "isFinanceOfficer() = true\nisFaculty() = true", "finance_hub.view: PASS\ncourses.view: PASS", "100% SUCCESS: User manages fee collections and PRN reconciliations in Bursary Hub, while accessing LMS course tools to deliver Accounting lectures."],
        ["Scenario 4:\nEstates + Dean", "Primary: Estates Manager\nSecondary: Dean", "isFacilitiesStaff() = true\nisDean() = true", "facilities.view: PASS\ndean_administration.view: PASS", "100% SUCCESS: User coordinates campus maintenance and hostel allocations, while presiding over Faculty Board budgets and strategic initiatives."]
    ]
    add_styled_table(doc, test_headers, test_data, [Inches(1.4), Inches(1.4), Inches(1.3), Inches(1.4), Inches(1.3)])

    add_image_with_caption(doc, "diagram_use_case_multirole_academic.png", "Multi-Role Academic Governance Use Case Diagram (Lecturer / HOD / Dean)")
    add_image_with_caption(doc, "diagram_sequence_multirole_auth.png", "Sequence Diagram: Multi-Role Authorization & Permission Resolution Pipeline")

    print("Part 1 built successfully!")
