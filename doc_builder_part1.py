"""
Comprehensive Document Generator for JBI University System Documentation (.docx)
Builds an exhaustive, highly detailed engineering & institutional governance specification.
"""

import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOC_OUTPUT_DIR = r"C:\Users\DIGIVERSE SOLUTIONS\Desktop\jbidocs"
IMAGE_DIR = os.path.join(DOC_OUTPUT_DIR, "images")
OUTPUT_DOCX = os.path.join(DOC_OUTPUT_DIR, "JBI_University_System_Architecture_and_Engineering_Documentation.docx")
SCHEMA_JSON = r"e:\Projects\JBI\storage\schema_dump.json"

os.makedirs(DOC_OUTPUT_DIR, exist_ok=True)

# Color Palette
COLOR_NAVY = RGBColor(0x1E, 0x3A, 0x8A)      # Primary Title & H1
COLOR_BLUE = RGBColor(0x25, 0x63, 0xEB)      # H2 Subheading
COLOR_SLATE = RGBColor(0x33, 0x41, 0x55)     # H3 Subheading
COLOR_BODY = RGBColor(0x1E, 0x29, 0x3B)      # Body text
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)     # Captions / Subtext
COLOR_AMBER = RGBColor(0xB4, 0x53, 0x09)     # Highlights / Warnings
COLOR_GREEN = RGBColor(0x15, 0x80, 0x3D)     # Success / Certified

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_padding(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLUE
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_SLATE
    return p

def add_heading_4(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = COLOR_BODY
    return p

def add_paragraph(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_NAVY
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10)
    run.font.italic = italic
    run.font.color.rgb = COLOR_BODY
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_NAVY
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_BODY
    return p

def add_callout(doc, title, text, bg_hex='F0FDF4', border_hex='16A34A', title_color=RGBColor(0x16, 0x65, 0x34)):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_padding(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"■ {title}\n")
    r1.font.name = 'Segoe UI'
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = title_color

    r2 = p.add_run(text)
    r2.font.name = 'Segoe UI'
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = COLOR_BODY

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)

def add_image_with_caption(doc, image_name, caption_text, width=Inches(6.4)):
    img_path = os.path.join(IMAGE_DIR, image_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run(f"Figure: {caption_text}")
        r_cap.font.name = 'Segoe UI'
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED
    else:
        p_warn = doc.add_paragraph(f"[Image Missing: {image_name}]")
        p_warn.paragraph_format.space_after = Pt(8)

def add_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], '1E3A8A')
        set_cell_padding(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_hex = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_hex)
            set_cell_padding(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_BODY

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(8)

print("Helper functions defined successfully!")
