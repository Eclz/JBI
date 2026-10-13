"""
Master Document Assembler for JBI University System Documentation
Assembles Part 1, Part 2, and Part 3 into the final .docx file on Desktop and in project docs.
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from doc_builder_part1 import DOC_OUTPUT_DIR, OUTPUT_DOCX
from doc_sections_part1 import build_part_1
from doc_sections_part2 import build_part_2
from doc_sections_part3 import build_part_3

def assemble_master_document():
    print("=" * 60)
    print("STARTING COMPLETE MASTER DOCUMENT GENERATION")
    print("=" * 60)

    doc = docx.Document()

    # Configure Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Configure Header & Footer
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.text = "JBI University Management System | Architecture & Engineering Specification"
        p_hdr.paragraph_format.space_after = Pt(4)
        for r in p_hdr.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.text = "Confidential - For University Board, Executive Leadership & Engineering Team Only"
        for r in p_ftr.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Build Sections
    print("\n[Step 1/3] Assembling Part 1 (Executive Summary, Actors, Multi-Role Architecture)...")
    build_part_1(doc)

    doc.add_page_break()

    print("\n[Step 2/3] Assembling Part 2 (Master Use Cases, Sequence Traces, DFD Levels 0/1/2)...")
    build_part_2(doc)

    doc.add_page_break()

    print("\n[Step 3/3] Assembling Part 3 (Data Dictionary, Normalization Proofs, Subsystems, Dev Reference, DevOps)...")
    build_part_3(doc)

    # Save Master Document
    print(f"\nSaving Master Document to: {OUTPUT_DOCX}")
    doc.save(OUTPUT_DOCX)

    # Also copy to local project directory e:\Projects\JBI\docs\
    project_docs_dir = r"e:\Projects\JBI\docs"
    os.makedirs(project_docs_dir, exist_ok=True)
    project_docx_path = os.path.join(project_docs_dir, "JBI_University_System_Architecture_and_Engineering_Documentation.docx")
    shutil.copy2(OUTPUT_DOCX, project_docx_path)
    print(f"Copied backup to: {project_docx_path}")

    # Inspect generated document statistics
    file_size_mb = os.path.getsize(OUTPUT_DOCX) / (1024 * 1024)
    total_paragraphs = len(doc.paragraphs)
    total_tables = len(doc.tables)
    
    # Calculate word count
    word_count = sum(len(p.text.split()) for p in doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                word_count += sum(len(p.text.split()) for p in cell.paragraphs)

    print("\n" + "=" * 60)
    print("MASTER DOCUMENT GENERATION COMPLETE!")
    print(f"File Path:        {OUTPUT_DOCX}")
    print(f"File Size:        {file_size_mb:.2f} MB")
    print(f"Total Paragraphs: {total_paragraphs}")
    print(f"Total Tables:     {total_tables}")
    print(f"Total Words:      ~{word_count:,} words")
    print("=" * 60)

if __name__ == "__main__":
    assemble_master_document()
