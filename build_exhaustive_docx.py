"""
Enhanced Word Document Builder to append full field-by-field schema tables for all 103 tables.
"""

import os
import json
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from ieee_data_definitions import (
    BASE_DIR, DOCS_DIR, DESKTOP_DIR,
    DOCX_OUTPUT_PATH, DESKTOP_DOCX_PATH,
    TABLE_DOMAINS, TABLE_DESCRIPTIONS, tables_dict
)

from scratch_build_ieee_spec import (
    add_heading_1, add_heading_2, add_heading_3, add_paragraph,
    add_callout, add_styled_table, set_cell_background, set_cell_padding,
    C_NAVY, C_BLUE, C_SLATE, C_BODY
)

def build_exhaustive_docx():
    print("\n>>> Compiling Complete Exhaustive IEEE Word (.docx) Document (with full 103-table field schemas)...")
    doc = docx.Document()

    # Configure 1.0 inch margins all around
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

        # Running Header
        hdr = s.header
        p_hdr = hdr.paragraphs[0]
        p_hdr.text = "JBI University Management System | IEEE Std 830-1998 & 1016-2009 Master Specification"
        p_hdr.paragraph_format.space_after = Pt(4)
        for r in p_hdr.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

        # Running Footer
        ftr = s.footer
        p_ftr = ftr.paragraphs[0]
        p_ftr.text = "Confidential — For Academic Board, Executive Directorate & Lead Engineering Team Only"
        for r in p_ftr.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Title Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("JBI UNIVERSITY MANAGEMENT SYSTEM")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = C_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Software Requirements Specification (SRS) & Software Design Description (SDD)\nConforming to ANSI/IEEE Std 830-1998, ISO/IEC/IEEE 29148:2018 & IEEE Std 1016-2009")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = C_BLUE

    # Meta Table
    meta_headers = ["Engineering Parameter", "Institutional Specification Details"]
    meta_rows = [
        ["System Reference", "JBI Tertiary Education Enterprise ERP (Academic, Financial & Governance)"],
        ["Target Audience", "University Board, Academic Senate, Faculty Deans, System Architects & DevOps"],
        ["Version & Release", "Version 2.2 Enterprise Production Release (October 2026)"],
        ["Framework & Stack", "Laravel 11.x, PHP 8.2+, MySQL 8.0 InnoDB, Bootstrap 5.3, Redis"],
        ["Applicable IEEE Standards", "IEEE Std 830-1998, ISO/IEC/IEEE 29148:2018, IEEE Std 1016-2009, IEEE Std 1012-2012"],
        ["Classification", "Institutional Confidential — Restricted to Board & Lead Engineering Team"],
        ["Document Identifier", "JBI-ENG-IEEE-830-1016-2026-V2.2"]
    ]
    add_styled_table(doc, meta_headers, meta_rows, [Inches(2.3), Inches(4.2)])

    add_callout(
        doc,
        "This master specification provides the formal mathematical, behavioral, relational, and architectural "
        "blueprints of the JBI University Management System ERP. All colloquial summaries are strictly deprecated. "
        "Every subsystem is formally specified with input contracts, algorithmic transformations, database integrity constraints, "
        "and verification methods.",
        "FORMAL IEEE ENGINEERING MANDATE"
    )

    doc.add_page_break()

    # Section 1: Introduction
    add_heading_1(doc, "1. Introduction (IEEE Std 830-1998 Section 1)")
    add_heading_2(doc, "1.1. Purpose & Scope")
    add_paragraph(doc, "This document establishes the definitive engineering baseline for the JBI University Management System ERP across its complete academic and operational lifecycle. The software encompasses ten major enterprise domains: Academic Affairs, Admissions, Assessments & LMS, Faculty Deanery, Finance & General Ledger, Human Resources, Campus Logistics & Hostels, E-Voting Democracy, Institutional Communications, and Kernel RBAC Security.")

    add_heading_2(doc, "1.2. Definitions, Acronyms, and Abbreviations")
    acronym_headers = ["Acronym / Term", "Formal Definition / Context"]
    acronym_data = [
        ["SRS", "Software Requirements Specification conforming to ANSI/IEEE Std 830-1998."],
        ["SDD", "Software Design Description conforming to IEEE Std 1016-2009."],
        ["RBAC", "Role-Based Access Control model enforced via many-to-many pivot relationship (role_user)."],
        ["PRN", "Payment Reference Number generated algorithmically for automated bank reconciliation."],
        ["GPA / CGPA", "Grade Point Average / Cumulative Grade Point Average calculated on institutional standard scale."],
        ["RTM", "Requirements Traceability Matrix establishing bidirectional mapping between SRS and code."],
        ["ACID", "Atomicity, Consistency, Isolation, and Durability transaction guarantees enforced by MySQL InnoDB."],
        ["MTBF", "Mean Time Between Failures (Quantitative reliability metric >= 720 operating hours)."],
        ["MTTR", "Mean Time To Recovery (Quantitative maintainability metric <= 15 minutes)."],
        ["HOD", "Head of Academic Department."],
        ["BCNF", "Boyce-Codd Normal Form satisfied across all relational schema tables."],
        ["V&V", "Verification and Validation conforming to IEEE Std 1012-2012."]
    ]
    add_styled_table(doc, acronym_headers, acronym_data, [Inches(1.8), Inches(4.7)])

    # Section 2: Overall Description
    add_heading_1(doc, "2. Overall System Description (IEEE Std 830-1998 Section 2)")
    add_heading_2(doc, "2.1. Multi-Tier Stakeholder Matrix")
    stakeholder_headers = ["Actor / Role", "Clearance", "Institutional Responsibilities", "Interface Scope"]
    stakeholder_data = [
        ["Super Administrator", "Tier-0", "Infrastructure maintenance, global settings, migrations, audit logs.", "/admin/*, Kernel"],
        ["University Registrar", "Tier-1", "Admissions, student matriculation, academic calendar, transcripts.", "/admin/students/*"],
        ["Faculty Dean", "Tier-2", "Collegiate oversight, budget approvals, quality reviews, petitions.", "/faculty/dean/*"],
        ["Department Head (HOD)", "Tier-3", "Curriculum scheduling, lecturer allocation, grade moderation.", "/faculty/department/*"],
        ["Academic Lecturer", "Tier-4", "Course syllabi, lecture notes, continuous assessment, exams.", "/faculty/courses/*"],
        ["Enrolled Student", "Tier-5", "Course registration, assignment submissions, exams, fees, voting.", "/student/*"],
        ["University Bursar", "Tier-2", "Fee structures, PRN reconciliation, budget disbursements, ledger.", "/admin/finance/*"],
        ["Human Resource Director", "Tier-2", "Staff recruitment, onboarding, leave approvals, payroll registers.", "/admin/hr/*"],
        ["Electoral Commissioner", "Tier-3", "Guild election setup, candidate vetting, certified ballot tallying.", "/admin/voting/*"]
    ]
    add_styled_table(doc, stakeholder_headers, stakeholder_data, [Inches(1.6), Inches(0.9), Inches(2.4), Inches(1.6)])

    # Section 3: Functional Requirements
    add_heading_1(doc, "3. System Functional Requirements Specification (FR-01 to FR-12)")
    add_paragraph(doc, "Each institutional functional requirement is formally specified with preconditions, input contracts, algorithmic state transformations, postconditions, and exception handling paths.")

    fr_docx_summary = [
        ["FR-01", "Multi-Role RBAC", "users, roles, role_user", "Authenticates credentials via bcrypt; dynamically resolves multi-role permissions; enforces RoleMiddleware."],
        ["FR-02", "Curriculum Hierarchy", "faculties, departments, courses", "Enforces 4-tier academic hierarchy; validates course codes and prerequisite linkages."],
        ["FR-03", "Admissions Pipeline", "applications, student_profiles", "Processes digital dossiers; assigns unique matriculation numbers; provisions student credentials."],
        ["FR-04", "Course Registration", "course_enrollments, courses", "Validates semester fee clearance; verifies passed prerequisites; enforces credit load bounds [15, 24]."],
        ["FR-05", "LMS & Coursework", "course_materials, assignments", "Manages lecture material distribution, assignment submissions, and timed online quizzes."],
        ["FR-06", "Grading & CGPA Engine", "grades, exams, exam_attempts", "Computes total scores (CA+Exam); maps to grade points; recalculates semester GPA and CGPA."],
        ["FR-07", "Deanery Governance", "budget_requests, quality_reviews", "Empowers Deans to arbitrate student petitions, evaluate faculty, and approve departmental budgets."],
        ["FR-08", "Fee Billing & PRN", "fee_structures, payments, ledger", "Generates algorithmic 12-digit PRN; reconciles bank payments; posts double-entry ledger journals."],
        ["FR-09", "HR & Payroll Engine", "hr_employees, payroll_records", "Manages staff contracts, biometric time logs, leave approvals, and automated monthly payroll registers."],
        ["FR-10", "Hostels & Logistics", "hostel_allocations, library_loans", "Allocates hostel beds by gender and capacity; manages library item circulation and asset tracking."],
        ["FR-11", "Campus E-Voting", "voting_sessions, votes", "Conducts secret ballot guild elections with cryptographic double-vote prevention."],
        ["FR-12", "Audit Kernel", "audit_logs, system_settings", "Maintains immutable system modification audit trail with client IP and serialized attributes."]
    ]
    add_styled_table(doc, ["Req ID", "Domain Title", "Primary Tables", "Algorithmic & Business Summary"], fr_docx_summary, [Inches(0.8), Inches(1.5), Inches(1.8), Inches(2.4)])

    doc.add_page_break()

    # Section 4: Non-Functional Requirements
    add_heading_1(doc, "4. Non-Functional Requirements & Performance Attributes")
    nfr_headers = ["Engineering Dimension", "Quantitative Institutional Specification", "Verification Standard"]
    nfr_data = [
        ["Read Latency Budget", "<= 200 ms for standard dashboard queries at 95th percentile.", "Automated Profiler Benchmark"],
        ["Write Latency Budget", "<= 500 ms for transactional mutations (enrollment, grading, PRN).", "JMeter Load Test Suite"],
        ["Database Query Budget", "<= 15 queries per HTTP request via eager loading (with()).", "Query Detector Inspection"],
        ["System Availability Target", "99.9% uptime during operational academic terms.", "Cloud Monitoring SLA"],
        ["Mean Time Between Failures", ">= 720 continuous operating hours (MTBF).", "Incident Register Log"],
        ["Mean Time To Recovery", "<= 15 minutes via automated container restart (MTTR).", "Disaster Recovery Drill"],
        ["Cryptographic Security", "Bcrypt work factor >= 12; TLS 1.3 encryption; prepared statements.", "Penetration Audit"]
    ]
    add_styled_table(doc, nfr_headers, nfr_data, [Inches(1.8), Inches(3.0), Inches(1.7)])

    # Section 5: Exhaustive Data Dictionary (All 103 tables with column-level schemas)
    add_heading_1(doc, "5. Exhaustive Relational Data Dictionary (All 103 Tables)")
    add_paragraph(doc, "Below is the complete database dictionary for all 103 relational tables operating under MySQL InnoDB engine, including column definitions, data types, nullability, keys, default values, and domain responsibilities.")

    for domain_name, tbl_list in TABLE_DOMAINS.items():
        add_heading_2(doc, domain_name)
        for t_name in tbl_list:
            t_cols = tables_dict.get(t_name, [])
            desc = TABLE_DESCRIPTIONS.get(t_name, "Domain entity table.")
            add_heading_3(doc, f"Table: {t_name} ({len(t_cols)} columns)")
            add_paragraph(doc, desc, bold_prefix="Domain Responsibility: ")

            col_headers = ["Field Name", "Data Type", "Null", "Key", "Default", "Description"]
            col_rows = []
            for c in t_cols:
                f_name = c.get("field", "")
                f_type = c.get("type", "")
                f_null = c.get("null", "")
                f_key = c.get("key", "")
                f_def = c.get("default")
                f_def_str = "NULL" if f_def is None else str(f_def)
                k_label = "PK" if f_key == "PRI" else ("FK/Index" if f_key in ["MUL", "UNI"] else "-")
                c_desc = "Primary Key" if f_name == "id" else ("Foreign Key" if f_name.endswith("_id") else ("Timestamps" if "at" in f_name else f_name.replace('_', ' ').capitalize()))
                col_rows.append([f_name, f_type, f_null, k_label, f_def_str, c_desc])

            add_styled_table(doc, col_headers, col_rows, [Inches(1.5), Inches(1.4), Inches(0.5), Inches(0.7), Inches(0.8), Inches(1.6)])

    doc.add_page_break()

    # Section 6: Formal Normalization Proofs
    add_heading_1(doc, "6. Formal Relational Database Normalization Proofs")
    add_heading_2(doc, "6.1. First Normal Form (1NF) Compliance Proof")
    add_paragraph(doc, "All 103 tables satisfy 1NF. Every column contains strictly atomic, indivisible values. Repeating groups and delimited arrays are eliminated through normalized join tables such as course_enrollments, role_user, and hr_shift_assignments.")

    add_heading_2(doc, "6.2. Second Normal Form (2NF) Compliance Proof")
    add_paragraph(doc, "All tables satisfy 2NF. In every composite key relation (e.g., role_user(user_id, role_id)), all non-key attributes are fully functionally dependent on the complete primary key, eliminating partial dependencies.")

    add_heading_2(doc, "6.3. Third Normal Form (3NF) Compliance Proof")
    add_paragraph(doc, "All tables satisfy 3NF. Transitive dependencies are decoupled through foreign key relationships. In students, collegiate details are not duplicated; department_id references departments, which in turn references faculties, preventing transitive update anomalies.")

    # Section 7: Requirements Traceability Matrix (RTM)
    add_heading_1(doc, "7. Requirements Traceability Matrix (RTM)")
    rtm_headers = ["SRS ID", "Functional Title", "Implementing Controller(s)", "Database Model / Table", "IEEE 1012 Method"]
    rtm_data = [
        ["FR-01", "Multi-Role RBAC", "LoginController, RoleMiddleware", "User, Role, role_user", "Test & Inspection"],
        ["FR-02", "Curriculum Setup", "CourseController, DepartmentController", "Faculty, Department, Course", "Test & Demonstration"],
        ["FR-03", "Admissions Pipeline", "ApplicationController, StudentController", "Application, StudentProfile", "Analysis & Demo"],
        ["FR-04", "Course Registration", "CourseRegistrationController", "CourseEnrollment, Course", "Test & Analysis"],
        ["FR-05", "LMS & Coursework", "AssignmentController, QuizController", "Assignment, QuizAttempt", "Test & Demonstration"],
        ["FR-06", "Grading & CGPA", "GradeController, ExamController", "Grade, ExamAttempt", "Test & Inspection"],
        ["FR-07", "Deanery Governance", "DeanController, BudgetController", "FacultyProfile, BudgetRequest", "Demo & Inspection"],
        ["FR-08", "Fee Billing & PRN", "FinanceController, PaymentController", "FeeRecord, Payment, GL_Entry", "Test & Audit"],
        ["FR-09", "HR & Payroll", "HrEmployeeController, PayrollController", "HrEmployee, PayrollRecord", "Test & Demo"],
        ["FR-10", "Hostels & Logistics", "HostelController, LibraryController", "HostelAllocation, LibraryLoan", "Test & Demo"],
        ["FR-11", "Campus E-Voting", "VotingController, ElectionController", "VotingSession, Vote", "Security Test"],
        ["FR-12", "Audit Kernel", "AuditLogController, SettingsController", "AuditLog, SystemSetting", "Inspection & Analysis"]
    ]
    add_styled_table(doc, rtm_headers, rtm_data, [Inches(0.8), Inches(1.5), Inches(1.8), Inches(1.4), Inches(1.0)])

    # Save Word Document
    doc.save(DOCX_OUTPUT_PATH)
    print(f"Exhaustive Word document saved to: {DOCX_OUTPUT_PATH}")

    # Copy to Desktop
    shutil.copy2(DOCX_OUTPUT_PATH, DESKTOP_DOCX_PATH)
    print(f"Exhaustive Word document copied to Desktop: {DESKTOP_DOCX_PATH}")

    file_size_mb = os.path.getsize(DOCX_OUTPUT_PATH) / (1024 * 1024)
    print(f"Exhaustive DOCX size: {file_size_mb:.2f} MB")

if __name__ == "__main__":
    build_exhaustive_docx()
